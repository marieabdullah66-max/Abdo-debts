from datetime import date
import csv
import io
import re
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, UploadFile
from openpyxl import load_workbook
from ..core import *
from ..xls_biff import read_first_sheet_rows

router = APIRouter(prefix="/api/suppliers", tags=["suppliers"])

MAX_SUPPLIER_IMPORT_BYTES = 12 * 1024 * 1024
MAX_SUPPLIER_IMPORT_ROWS = 5000
MAX_SUPPLIER_IMPORT_COLS = 40


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").replace("\ufeff", "").strip().split())


def _parse_money(value: Any) -> float:
    raw = _clean_text(value).replace(",", "")
    if not raw:
        return 0.0
    # Preserve the sign from the external report:
    # positive = debt for us, negative = debt on us, zero = no debt.
    negative = False
    if raw.endswith("-"):
        negative = True
        raw = raw[:-1].strip()
    if raw.startswith("(") and raw.endswith(")"):
        negative = True
        raw = raw[1:-1].strip()
    cleaned = re.sub(r"[^0-9.\-]", "", raw)
    if cleaned.count("-") > 1:
        cleaned = cleaned.replace("-", "")
        negative = True
    if cleaned.startswith("-"):
        negative = True
        cleaned = cleaned[1:]
    if cleaned in {"", ".", "-"}:
        return 0.0
    try:
        amount = float(Decimal(cleaned))
    except (InvalidOperation, ValueError):
        return 0.0
    return round(-abs(amount) if negative else amount, 2)


def _debt_state(value: float) -> str:
    amount = round(float(value or 0), 2)
    if amount > 0:
        return "debt_for_us"
    if amount < 0:
        return "debt_on_us"
    return "no_debt"


def _parse_date(value: Any) -> date | None:
    raw = _clean_text(value)
    if not raw or raw in {".", "-", "غيرمحدد"}:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            from datetime import datetime
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def _looks_like_date(value: str) -> bool:
    return _parse_date(value) is not None


def _supplier_import_row(cells: list[Any], idx: int) -> dict[str, Any] | None:
    cells = [_clean_text(x) for x in cells]
    if not cells:
        return None

    # CSV export: labels repeat inside every row. Useful cells are fixed.
    if len(cells) > 17 and (cells[6] == "الرصيد" or cells[7] == "اسم الزبون" or cells[8] == "ر.م"):
        name = cells[13] if len(cells) > 13 else ""
        reference_no = cells[14] if len(cells) > 14 else ""
        balance = _parse_money(cells[12] if len(cells) > 12 else "")
        last_payment_date = _parse_date(cells[16] if len(cells) > 16 else "")
        last_invoice_date = _parse_date(cells[17] if len(cells) > 17 else "")
    # XLS export from Crystal Reports: first visible columns are labels, then
    # data appears as: ... الرصيد, اسم الزبون, ر.م, ..., تاريخ آخر سداد, تاريخ آخر فاتورة.
    elif len(cells) >= 6:
        name = cells[4] if len(cells) > 4 else ""
        reference_no = cells[5] if len(cells) > 5 else ""
        balance = _parse_money(cells[3] if len(cells) > 3 else "")
        last_payment_date = _parse_date(cells[7] if len(cells) > 7 else "")
        last_invoice_date = _parse_date(cells[8] if len(cells) > 8 else "")
    else:
        return None

    # Fallback: choose a likely supplier name from the later cells.
    if not name or name in {".", "غيرمحدد", "اسم الزبون"} or _looks_like_date(name) or re.fullmatch(r"[0-9.,\-]+", name or ""):
        for c in cells[4:]:
            if not c or c in {".", "غيرمحدد", "الاجمالي :", "اسم الزبون", "ر.م"}:
                continue
            if _looks_like_date(c) or re.fullmatch(r"[0-9.,\-]+", c):
                continue
            name = c
            break

    if not name or name in {".", "غيرمحدد", "اسم الزبون"}:
        return None

    return {
        "source_row": idx,
        "name": name[:160],
        "reference_no": reference_no[:100] or None,
        "balance": round(balance, 2),
        "debt_state": _debt_state(balance),
        "debt_label": "دين لنا" if balance > 0 else ("دين علينا" if balance < 0 else "بدون دين"),
        "last_payment_date": last_payment_date.isoformat() if last_payment_date else None,
        "last_invoice_date": last_invoice_date.isoformat() if last_invoice_date else None,
        "include": True,
    }


def _parse_supplier_rows(source_rows: list[list[Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for idx, row in enumerate(source_rows, start=1):
        parsed = _supplier_import_row(row, idx)
        if parsed:
            rows.append(parsed)
    return rows


def _parse_supplier_csv_rows(text: str) -> list[dict[str, Any]]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except csv.Error:
        dialect = csv.excel
    return _parse_supplier_rows(list(csv.reader(io.StringIO(text), dialect)))


def _parse_supplier_xlsx_rows(raw: bytes) -> list[dict[str, Any]]:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        sheet_rows = []
        for row in ws.iter_rows(max_row=MAX_SUPPLIER_IMPORT_ROWS + 10, max_col=MAX_SUPPLIER_IMPORT_COLS, values_only=True):
            values = list(row or [])
            while values and (values[-1] is None or _clean_text(values[-1]) == ""):
                values.pop()
            sheet_rows.append(values)
        return _parse_supplier_rows(sheet_rows)
    except Exception as exc:
        raise HTTPException(422, "تعذر قراءة ملف .xlsx") from exc


def _parse_supplier_xls_rows(raw: bytes) -> list[dict[str, Any]]:
    try:
        source_rows = read_first_sheet_rows(raw, max_rows=MAX_SUPPLIER_IMPORT_ROWS + 10, max_cols=MAX_SUPPLIER_IMPORT_COLS)
        return _parse_supplier_rows(source_rows)
    except Exception as exc:
        raise HTTPException(422, "تعذر قراءة ملف .xls") from exc


async def _unique_import_invoice_number(supplier_id: str, base: str) -> str:
    base = _clean_text(base)[:90] or "IMPORT"
    existing = await sb("GET", "/rest/v1/invoices", service=True, params={"select": "invoice_number", "supplier_id": f"eq.{supplier_id}", "invoice_number": f"ilike.{base}%", "limit": "10000"})
    used = {_clean_text(x.get("invoice_number")) for x in existing or []}
    if base not in used:
        return base
    for i in range(2, 10000):
        candidate = f"{base}-{i}"
        if candidate not in used:
            return candidate
    return f"{base}-{int(time.time())}"

async def _external_balances(branch_id: str | None, profile: dict[str, Any]) -> list[dict[str, Any]]:
    params: dict[str, str] = {"select": "supplier_id,branch_id,signed_balance,last_invoice_date,last_payment_date,reference_no", "limit": "10000"}
    params = apply_branch_filter(params, profile)
    if branch_id:
        require_branch_access(profile, branch_id)
        params["branch_id"] = f"eq.{branch_id}"
    try:
        return await sb("GET", "/rest/v1/supplier_external_balances", service=True, params=params) or []
    except HTTPException as exc:
        detail = str(getattr(exc, "detail", ""))
        if "supplier_external_balances" in detail or "does not exist" in detail:
            return []
        raise


async def _upsert_external_balance(*, supplier_id: str, branch_id: str, signed_balance: float, row: SupplierImportRowInput, profile: dict[str, Any]) -> None:
    payload = {
        "supplier_id": supplier_id,
        "branch_id": branch_id,
        "signed_balance": round(float(signed_balance or 0), 2),
        "reference_no": row.reference_no,
        "last_payment_date": row.last_payment_date.isoformat() if row.last_payment_date else None,
        "last_invoice_date": row.last_invoice_date.isoformat() if row.last_invoice_date else None,
        "updated_by": profile["id"],
    }
    try:
        await sb(
            "POST",
            "/rest/v1/supplier_external_balances",
            service=True,
            headers={"Prefer": "resolution=merge-duplicates"},
            params={"on_conflict": "supplier_id,branch_id"},
            json=payload,
        )
    except HTTPException as exc:
        detail = str(getattr(exc, "detail", ""))
        if "supplier_external_balances" in detail or "does not exist" in detail:
            raise HTTPException(500, "يلزم تشغيل SQL الخاص بإصدار V65 قبل استيراد أرصدة الموردين بالإشارة") from exc
        raise


async def _category_map() -> dict[str, list[dict[str, Any]]]:
    links = await sb(
        "GET", "/rest/v1/supplier_category_links", service=True,
        params={
            "select": "supplier_id,category_id,supplier_categories(id,name)",
            "limit": "20000",
        },
    )
    result: dict[str, list[dict[str, Any]]] = {}
    for link in links or []:
        supplier_id = link.get("supplier_id")
        category = link.get("supplier_categories") or {}
        if not supplier_id or not category.get("id"):
            continue
        result.setdefault(supplier_id, []).append({"id": category.get("id"), "name": category.get("name") or ""})
    for rows in result.values():
        rows.sort(key=lambda x: (x.get("name") or "").lower())
    return result


async def _validate_category_ids(category_ids: list[str]) -> list[str]:
    ids = list(dict.fromkeys(category_ids or []))
    if not ids:
        return []
    rows = await sb(
        "GET", "/rest/v1/supplier_categories", service=True,
        params={"select": "id", "id": f"in.({','.join(ids)})", "limit": str(max(100, len(ids)))},
    )
    found = {row.get("id") for row in rows or []}
    if found != set(ids):
        raise HTTPException(422, "يوجد تصنيف مورد غير صالح")
    return ids


async def _replace_supplier_categories(supplier_id: str, category_ids: list[str]) -> None:
    await sb("DELETE", "/rest/v1/supplier_category_links", service=True, params={"supplier_id": f"eq.{supplier_id}"})
    if category_ids:
        await sb(
            "POST", "/rest/v1/supplier_category_links", service=True,
            json=[{"supplier_id": supplier_id, "category_id": category_id} for category_id in category_ids],
        )


@router.post("/import/preview")
async def preview_supplier_import(file: UploadFile = File(...), profile: dict[str, Any] = Depends(current_profile)) -> Any:
    require_permission(profile, "manage_suppliers")
    raw = await file.read(MAX_SUPPLIER_IMPORT_BYTES + 1)
    if len(raw) > MAX_SUPPLIER_IMPORT_BYTES:
        raise HTTPException(413, "حجم الملف كبير جدًا؛ الحد الأقصى 12 MB")
    if not raw:
        raise HTTPException(422, "الملف فارغ")
    name = (file.filename or "").lower()
    content_type = (file.content_type or "").lower()
    if name.endswith(".xls"):
        rows = _parse_supplier_xls_rows(raw)
    elif name.endswith(".xlsx"):
        rows = _parse_supplier_xlsx_rows(raw)
    elif name.endswith(".csv") or content_type in {"text/csv", "application/csv", "application/vnd.ms-excel"}:
        decoded = None
        for enc in ("utf-8-sig", "utf-8", "cp1256", "windows-1256", "latin-1"):
            try:
                decoded = raw.decode(enc)
                break
            except UnicodeDecodeError:
                pass
        if decoded is None:
            raise HTTPException(422, "تعذر قراءة ملف CSV")
        rows = _parse_supplier_csv_rows(decoded)
    else:
        raise HTTPException(422, "الاستيراد يدعم CSV و Excel بصيغة .xls أو .xlsx")
    return {"rows": rows[:MAX_SUPPLIER_IMPORT_ROWS], "total_rows": len(rows), "filename": file.filename}


@router.post("/import")
async def import_suppliers(data: SupplierImportInput, profile: dict[str, Any] = Depends(current_profile)) -> Any:
    require_permission(profile, "manage_suppliers")
    require_permission(profile, "create_invoices")
    require_branch_access(profile, data.branch_id)
    branch = await sb("GET", "/rest/v1/branches", service=True, params={"select": "id,name", "id": f"eq.{data.branch_id}", "active": "eq.true", "limit": "1"})
    if not branch:
        raise HTTPException(422, "الفرع غير موجود أو موقوف")

    existing = await sb("GET", "/rest/v1/suppliers", service=True, params={"select": "id,name", "active": "eq.true", "limit": "10000"})
    by_name = {_clean_text(s.get("name")).casefold(): s for s in existing or []}
    created = 0
    reused = 0
    invoices_created = 0
    debt_for_us_count = 0
    debt_on_us_count = 0
    no_debt_count = 0
    skipped = 0
    imported_rows = []

    for idx, row in enumerate(data.rows, start=1):
        if not row.include:
            skipped += 1
            continue
        clean_name = _clean_text(row.name)
        if len(clean_name) < 2:
            skipped += 1
            continue
        key = clean_name.casefold()
        supplier = by_name.get(key)
        if not supplier:
            inserted = await sb("POST", "/rest/v1/suppliers", service=True, headers={"Prefer": "return=representation"}, json={
                "name": clean_name,
                "phone": None,
                "notes": f"استيراد خارجي - ر.م: {row.reference_no}" if row.reference_no else "استيراد خارجي",
                "created_by": profile["id"],
                "active": True,
            })
            supplier = inserted[0]
            by_name[key] = supplier
            created += 1
        else:
            reused += 1
            if row.reference_no:
                await sb("PATCH", "/rest/v1/suppliers", service=True, params={"id": f"eq.{supplier['id']}"}, json={"notes": f"استيراد خارجي - ر.م: {row.reference_no}"})

        signed_balance = round(float(row.balance or 0), 2)
        if signed_balance > 0:
            debt_for_us_count += 1
        elif signed_balance < 0:
            debt_on_us_count += 1
        else:
            no_debt_count += 1
        await _upsert_external_balance(supplier_id=supplier["id"], branch_id=data.branch_id, signed_balance=signed_balance, row=row, profile=profile)

        # In the external supplier report: negative value = debt on us.
        # Keep the ordinary invoice system positive by creating an opening invoice
        # only for values we owe, using the absolute amount.
        if signed_balance < 0:
            invoice_date = row.last_invoice_date or date.today()
            ref = _clean_text(row.reference_no) or str(idx)
            invoice_number = await _unique_import_invoice_number(supplier["id"], f"IMPORT-{ref}-{invoice_date.isoformat()}")
            await sb("POST", "/rest/v1/invoices", service=True, headers={"Prefer": "return=representation"}, json={
                "supplier_id": supplier["id"],
                "branch_id": data.branch_id,
                "invoice_number": invoice_number,
                "amount": abs(signed_balance),
                "invoice_date": invoice_date.isoformat(),
                "due_date": None,
                "notes": f"رصيد مستورد من ملف خارجي - دين علينا{(' - آخر سداد: ' + row.last_payment_date.isoformat()) if row.last_payment_date else ''}",
                "created_by": profile["id"],
            })
            invoices_created += 1
        imported_rows.append({"name": clean_name, "supplier_id": supplier["id"], "balance": signed_balance, "debt_state": _debt_state(signed_balance)})

    return {
        "ok": True,
        "created_suppliers": created,
        "existing_suppliers": reused,
        "invoices_created": invoices_created,
        "debt_for_us_count": debt_for_us_count,
        "debt_on_us_count": debt_on_us_count,
        "no_debt_count": no_debt_count,
        "skipped": skipped,
        "rows": imported_rows[:50],
    }



@router.post("/reset-values")
async def reset_supplier_values(data: SupplierResetValuesInput, profile: dict[str, Any] = Depends(current_profile)) -> Any:
    """Delete financial values for one selected branch while keeping supplier names/master data.

    This is designed for periodic external debt-report imports per branch: old
    balances, imported opening invoices, payments, allocations, and payment
    plans are cleared only for the chosen branch, then a fresh supplier report
    can be imported for that branch without affecting other branches.
    """
    require_permission(profile, "manage_suppliers")
    require_branch_access(profile, data.branch_id)
    branch = await sb("GET", "/rest/v1/branches", service=True, params={"select": "id,name", "id": f"eq.{data.branch_id}", "active": "eq.true", "limit": "1"})
    if not branch:
        raise HTTPException(422, "الفرع غير موجود أو موقوف")

    branch_filter = {"branch_id": f"eq.{data.branch_id}", "limit": "10000"}
    existing_invoices = await sb("GET", "/rest/v1/invoices", service=True, params={"select": "id", **branch_filter})
    existing_payments = await sb("GET", "/rest/v1/payments", service=True, params={"select": "id", **branch_filter})
    existing_plans = await sb("GET", "/rest/v1/payment_plans", service=True, params={"select": "id", **branch_filter})

    invoice_ids = [row.get("id") for row in existing_invoices or [] if row.get("id")]
    payment_ids = [row.get("id") for row in existing_payments or [] if row.get("id")]

    def chunks(values: list[str], size: int = 80):
        for i in range(0, len(values), size):
            yield values[i:i + size]

    # Delete children/references for the selected branch first, then branch financial documents.
    await sb("DELETE", "/rest/v1/payment_plans", service=True, params={"branch_id": f"eq.{data.branch_id}"})
    for ids in chunks(payment_ids):
        await sb("DELETE", "/rest/v1/payment_allocations", service=True, params={"payment_id": f"in.({','.join(ids)})"})
    for ids in chunks(invoice_ids):
        await sb("DELETE", "/rest/v1/payment_allocations", service=True, params={"invoice_id": f"in.({','.join(ids)})"})
    await sb("DELETE", "/rest/v1/payments", service=True, params={"branch_id": f"eq.{data.branch_id}"})
    await sb("DELETE", "/rest/v1/invoices", service=True, params={"branch_id": f"eq.{data.branch_id}"})
    try:
        await sb("DELETE", "/rest/v1/supplier_external_balances", service=True, params={"branch_id": f"eq.{data.branch_id}"})
    except HTTPException as exc:
        detail = str(getattr(exc, "detail", ""))
        if "supplier_external_balances" not in detail and "does not exist" not in detail:
            raise

    return {
        "ok": True,
        "branch_id": data.branch_id,
        "branch_name": (branch[0] or {}).get("name"),
        "deleted_invoices": len(invoice_ids),
        "deleted_payments": len(payment_ids),
        "deleted_payment_plans": len(existing_plans or []),
    }


@router.get("")
async def list_suppliers(q: str | None = None, branch_id: str | None = None, include_balance: bool = False, profile: dict[str, Any] = Depends(current_profile)) -> Any:
    if not (effective_permissions(profile).get("view_suppliers") or effective_permissions(profile).get("view_payment_plans")):
        raise HTTPException(403, "ليس لديك صلاحية عرض الموردين")
    params = {"select": "id,name,phone,notes,active,created_at", "active": "eq.true", "order": "name.asc", "limit": "5000"}
    if q:
        safe = q.strip().replace("%", "")[:80]
        if safe:
            params["name"] = f"ilike.*{safe}*"
    suppliers = await sb("GET", "/rest/v1/suppliers", service=True, params=params)
    categories_by_supplier = await _category_map()

    if not include_balance and not branch_id:
        return [{**supplier, "categories": categories_by_supplier.get(supplier.get("id"), [])} for supplier in (suppliers or [])]

    inv_params: dict[str, str] = {"select": "supplier_id,balance,invoice_date", "limit": "10000"}
    inv_params = apply_branch_filter(inv_params, profile)
    if branch_id:
        require_branch_access(profile, branch_id)
        inv_params["branch_id"] = f"eq.{branch_id}"
    invoices = await sb("GET", "/rest/v1/invoice_balances", service=True, params=inv_params)
    external_rows = await _external_balances(branch_id, profile)
    external_by_supplier: dict[str, dict[str, Any]] = {row.get("supplier_id"): row for row in external_rows if row.get("supplier_id")}

    balances: dict[str, float] = {}
    oldest_open_invoice: dict[str, date] = {}
    suppliers_in_branch: set[str] = set()
    for inv in invoices or []:
        sid = inv.get("supplier_id")
        if not sid:
            continue
        suppliers_in_branch.add(sid)
        balance = float(inv.get("balance") or 0)
        balances[sid] = balances.get(sid, 0.0) + balance
        if balance <= 0 or not inv.get("invoice_date"):
            continue
        try:
            invoice_date = date.fromisoformat(str(inv.get("invoice_date")))
        except ValueError:
            continue
        previous = oldest_open_invoice.get(sid)
        if previous is None or invoice_date < previous:
            oldest_open_invoice[sid] = invoice_date

    suppliers_in_branch.update(external_by_supplier.keys())
    today = date.today()
    rows = []
    for supplier in suppliers or []:
        if branch_id and supplier.get("id") not in suppliers_in_branch:
            continue
        external = external_by_supplier.get(supplier.get("id")) or {}
        signed_balance = round(float(external.get("signed_balance") if external else balances.get(supplier.get("id"), 0.0) or 0.0), 2)
        rows.append({
            **supplier,
            "categories": categories_by_supplier.get(supplier.get("id"), []),
            "balance": signed_balance,
            "debt_state": _debt_state(signed_balance),
            "debt_label": "دين لنا" if signed_balance > 0 else ("دين علينا" if signed_balance < 0 else "بدون دين"),
            "import_reference_no": external.get("reference_no"),
            "aging_days": (
                max(0, (today - oldest_open_invoice[supplier.get("id")]).days)
                if supplier.get("id") in oldest_open_invoice else None
            ),
            "oldest_open_invoice_date": (
                oldest_open_invoice[supplier.get("id")].isoformat()
                if supplier.get("id") in oldest_open_invoice else None
            ),
        })
    return rows


@router.post("")
async def create_supplier(data: SupplierInput, profile: dict[str, Any] = Depends(current_profile)) -> Any:
    require_permission(profile, "manage_suppliers")
    category_ids = await _validate_category_ids(data.category_ids)
    rows = await sb("POST", "/rest/v1/suppliers", service=True, headers={"Prefer": "return=representation"}, json={
        "name": data.name.strip(), "phone": (data.phone or "").strip() or None, "notes": (data.notes or "").strip() or None,
        "created_by": profile["id"], "active": True,
    })
    supplier = rows[0]
    try:
        await _replace_supplier_categories(supplier["id"], category_ids)
    except Exception:
        await sb("DELETE", "/rest/v1/suppliers", service=True, params={"id": f"eq.{supplier['id']}"})
        raise
    supplier["categories"] = [x for x in (await _category_map()).get(supplier["id"], [])]
    return supplier


@router.put("/{supplier_id}")
async def update_supplier(supplier_id: str, data: SupplierInput, profile: dict[str, Any] = Depends(current_profile)) -> Any:
    require_permission(profile, "manage_suppliers")
    category_ids = await _validate_category_ids(data.category_ids)
    rows = await sb("PATCH", "/rest/v1/suppliers", service=True, headers={"Prefer": "return=representation"}, params={"id": f"eq.{supplier_id}"}, json={
        "name": data.name.strip(), "phone": (data.phone or "").strip() or None, "notes": (data.notes or "").strip() or None,
    })
    if not rows:
        raise HTTPException(404, "المورد غير موجود")
    await _replace_supplier_categories(supplier_id, category_ids)
    rows[0]["categories"] = (await _category_map()).get(supplier_id, [])
    return rows[0]


@router.delete("/{supplier_id}")
async def delete_supplier(supplier_id: str, profile: dict[str, Any] = Depends(current_profile)) -> Any:
    require_permission(profile, "manage_suppliers")
    linked = await sb("GET", "/rest/v1/invoices", service=True, params={"select": "id", "supplier_id": f"eq.{supplier_id}", "limit": "1"})
    if linked:
        raise HTTPException(409, "لا يمكن حذف المورد لوجود فواتير مرتبطة به")
    await sb("DELETE", "/rest/v1/suppliers", service=True, params={"id": f"eq.{supplier_id}"})
    return {"ok": True}


@router.get("/{supplier_id}/summary")
async def supplier_summary(supplier_id: str, profile: dict[str, Any] = Depends(current_profile)) -> Any:
    require_permission(profile, "view_suppliers")
    supplier = await sb("GET", "/rest/v1/suppliers", service=True, params={"select": "id,name,phone,notes", "id": f"eq.{supplier_id}", "limit": "1"})
    if not supplier:
        raise HTTPException(404, "المورد غير موجود")
    supplier[0]["categories"] = (await _category_map()).get(supplier_id, [])
    params = apply_branch_filter({
        "select": "id,invoice_number,amount,paid_amount,balance,status,invoice_date,due_date,notes,pdf_path,branch_id,branch_name",
        "supplier_id": f"eq.{supplier_id}", "order": "invoice_date.desc", "limit": "5000"
    }, profile)
    invoices = await sb("GET", "/rest/v1/invoice_balances", service=True, params=params)
    totals = {
        "invoiced": round(sum(float(x.get("amount") or 0) for x in invoices or []), 2),
        "paid": round(sum(float(x.get("paid_amount") or 0) for x in invoices or []), 2),
        "balance": round(sum(float(x.get("balance") or 0) for x in invoices or []), 2),
    }
    by_branch: dict[str, dict[str, Any]] = {}
    for inv in invoices or []:
        bid = inv.get("branch_id")
        row = by_branch.setdefault(bid, {"branch_id": bid, "branch_name": inv.get("branch_name", ""), "invoiced": 0.0, "paid": 0.0, "balance": 0.0})
        row["invoiced"] += float(inv.get("amount") or 0)
        row["paid"] += float(inv.get("paid_amount") or 0)
        row["balance"] += float(inv.get("balance") or 0)
    return {"supplier": supplier[0], "totals": totals, "by_branch": list(by_branch.values()), "invoices": [{**x, "branches": {"name": x.get("branch_name")}} for x in (invoices or [])]}
